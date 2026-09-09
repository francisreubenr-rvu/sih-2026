#!/usr/bin/env python3
"""Build the expanded working deck. This is not the six-slide submission.

Run with the bundled python runtime (python-pptx required).
All visuals are editable PowerPoint shapes. No external assets or fonts needed.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Docs' / 'expanded-pitch-blueprint.pptx'
PAPER, INK, LIME, FOREST, MUTED = 'F3F1E7', '15211F', 'D8F36A', '29473F', '52645C'
WHITE, CAUTION = 'FFFFFF', '70420B'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = 'SIH2171 | Expanded pitch blueprint | NOT SUBMISSION READY'
prs.core_properties.subject = '15-slide editable preparation blueprint; problem identity and domain remain unverified'
prs.core_properties.author = 'RV University team preparation workspace'
prs.core_properties.keywords = 'blueprint, unverified, SIH2171, not submission ready'

def col(v): return RGBColor.from_string(v)
def box(sl, x,y,w,h, fill=PAPER, line=None, shape=MSO_SHAPE.RECTANGLE):
    s=sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=col(fill)
    s.line.fill.background() if line is None else None
    if line: s.line.color.rgb=col(line)
    return s

def txt(sl,text,x,y,w,h,size=24,color=INK,bold=False,font='Arial'):
    sh=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=col(color)
        p.space_after=Pt(5)
    return sh

def rule(sl,x,y,w,color=MUTED):
    ln=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x+w), Inches(y))
    ln.line.color.rgb=col(color); ln.line.width=Pt(.7)

def arrow(sl,x,y,w=.44,color=FOREST):
    box(sl,x,y,w,.25,fill=color,shape=MSO_SHAPE.CHEVRON)

def panel(sl,x,y,w,h,label,body,status='EVIDENCE SLOT',dark=False):
    bg=INK if dark else WHITE; fg=PAPER if dark else INK
    box(sl,x,y,w,h,bg)
    txt(sl,status,x+.22,y+.18,w-.44,.28,11,LIME if dark else CAUTION,True)
    txt(sl,label,x+.22,y+.66,w-.44,.65,24,fg,True)
    txt(sl,body,x+.22,y+1.48,w-.44,h-1.6,18,PAPER if dark else MUTED)

def base(family,title,kicker,number,dark=False):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid(); sl.background.fill.fore_color.rgb=col(INK if dark else PAPER)
    fg=PAPER if dark else INK
    txt(sl,family,.55,.32,8,.26,11,LIME if dark else FOREST,True)
    box(sl,9.48,.23,3.3,.43,LIME)
    txt(sl,'NOT SUBMISSION READY',9.65,.32,3.0,.22,12,INK,True)
    txt(sl,title,.55,.94,12.1,1.45 if '\n' in title else 1.12,36 if '\n' in title else 39,fg,True)
    txt(sl,kicker,.58,2.48 if '\n' in title else 2.08,12.1,.6,18,PAPER if dark else MUTED)
    rule(sl,.55,6.96,12.23,FOREST if dark else MUTED)
    txt(sl,'SIH2171 · IDENTITY UNVERIFIED  /  EXPANDED BLUEPRINT',.55,7.12,10,.2,10,PAPER if dark else MUTED)
    txt(sl,f'{number:02d} / 15',11.9,7.12,.9,.2,10,PAPER if dark else MUTED)
    return sl

def notes(sl,seconds,text,refs=''):
    sl.notes_slide.notes_text_frame.text=(f'BLUEPRINT — NOT SUBMISSION READY\nPlanned slot: {seconds} seconds. Timing is an allocation, not a measured rehearsal.\n\n'+text+'\n\n'+('Sources / provenance:\n'+refs if refs else 'Provenance: authored planning framework; no measured outcome claimed.'))

F=['TITLE PAGE','IDEA TITLE','TECHNICAL APPROACH','FEASIBILITY AND VIABILITY','IMPACT AND BENEFITS','RESEARCH  AND REFERENCES']

s=base(F[0],'An entry begins with\na verified problem.','SIH 2026 preparation · RV University · Six-person second-year B.Tech CSE team',1,True)
box(s,.58,3.25,7.55,2.7,LIME)
txt(s,'SIH2171',.85,3.48,7,.8,56,INK,True)
txt(s,'Requested ID. Official title, theme and category remain unverified.',.87,4.57,6.7,1.0,23,INK)
txt(s,'Before this becomes a pitch',8.72,3.48,3.8,.5,22,PAPER,True)
txt(s,'Confirm statement identity\nConfirm registered team details\nReplace evidence slots',8.72,4.25,3.8,1.65,21,PAPER)
notes(s,25,'This file is an expanded preparation blueprint, not a completed competition pitch. The supplied identifier is SIH2171. Its official problem title, theme, category and edition have not been verified. The team facts come from Francis’s brief. Before presenting this as a solution, replace these open fields with the exact organizer-issued statement and registered team details. The separate submission deck must follow the supplied six-slide format.','User brief; Docs/team-plan.md; supplied SIH2025-IDEA-Presentation-Format (2).pptx')

s=base(F[1],'Make the problem measurable.','A population, a burden and a baseline — all three need sources.',2)
panel(s,.58,3.06,3.85,2.92,'WHO','Target audience\nGeography and context\nSource + date')
panel(s,4.72,3.06,3.85,2.92,'WHAT IT COSTS','Time, money or access\nUnit and denominator\nMeasurement method')
panel(s,8.86,3.06,3.85,2.92,'WHAT EXISTS','Present workflow\nBaseline result\nKnown limitations')
txt(s,'No domain statistic has been inserted. The problem must be verified first.',.6,6.35,12,.3,16,CAUTION)
notes(s,40,'Open the eventual pitch with one burden that matters to a defined user. Name the population, location and timeframe; explain the denominator behind any percentage. Use a primary source or a documented observation, and place its citation on the slide. Do not borrow a national statistic as proof of a local pain point without showing the connection. Today those domain fields are intentionally empty. The completion task is to verify the problem and establish a baseline, not to invent a compelling number. Once populated, state what changes for the user if the bottleneck is removed.')

s=base(F[1],'Locate the moment that breaks.','Map a real user journey before choosing a feature list.',3)
for x,n,t in [(0.58,'01','Need arises'),(3.72,'02','Current action'),(6.86,'03','Bottleneck'),(10.0,'04','Consequence')]:
    box(s,x,3.38,2.7,1.9,LIME if n=='03' else WHITE)
    txt(s,n,x+.2,3.55,2.3,.36,18,FOREST,True)
    txt(s,t,x+.2,4.15,2.3,.7,24,INK,True)
for x in [3.35,6.49,9.63]: arrow(s,x,4.21,.28)
txt(s,'Complete with an observed task, a source and the actual failure condition.',.6,5.79,11.9,.62,22,MUTED)
notes(s,40,'Use a single representative task from the verified domain. Walk from the initial need through the current action to the exact failure and its consequence. The four boxes are a research scaffold, not a discovered user journey. Support each material step with an interview record, operational source or observed task once available. Explain why the bottleneck matters rather than describing every screen. This slide should make the next design choice feel necessary. If the team cannot state where the workflow fails, return to the evidence before adding functionality.')

s=base(F[1],'A gap needs a fair comparison.','Compare the same task under the same conditions; unknown is a valid result.',4)
rows=[['APPROACH','CORE TASK','LIMITATION','SOURCE'],['Existing tool A','Unverified','Unverified','Required'],['Existing tool B','Unverified','Unverified','Required'],['Proposed approach','Not defined','Not tested','Future evidence']]
widths=[3.2,2.8,3.0,2.65]
for i,row in enumerate(rows):
    x=.58; y=3.06+i*.66
    for w,cell in zip(widths,row):
        box(s,x,y,w-.03,.62,INK if i==0 else WHITE)
        txt(s,cell,x+.16,y+.15,w-.28,.4,14 if i==0 else 18,PAPER if i==0 else INK,i==0)
        x+=w
txt(s,'Novelty claim: pending a sourced, domain-specific competitor review.',.6,6.22,12,.36,20,CAUTION)
notes(s,45,'Select actual competing approaches only after the domain is known. Include the incumbent manual workflow when it is a realistic alternative. Compare a small number of material dimensions: completion of the primary task, operating conditions, failure behavior and cost, for example. Verify claims in first-party documentation or by reproducible tests. Do not mark a competitor as lacking a feature because its website is silent. Unknown stays unknown. The future differentiation sentence must identify a demonstrated improvement and the conditions under which it holds. At this stage no competitor or novelty claim is established.')

s=base(F[1],'One promise. One mechanism.','Define the minimum complete solution after the problem identity is confirmed.',5,True)
txt(s,'FOR [VERIFIED USER]',.63,3.22,8,.4,15,LIME,True)
txt(s,'Enable [specific outcome]\nthrough [testable mechanism].',.63,3.86,11.8,1.75,36,PAPER,True)
txt(s,'Success condition → [measurable task result]     Boundary → [explicit exclusion]',.63,6.06,12,.5,19,PAPER)
notes(s,40,'Use this sentence to constrain the eventual prototype. Identify a verified user, the outcome they need and the implemented mechanism that produces it. Name a success condition that can be tested end to end. Add a boundary so judges understand the limits of the claim and the team has a defensible scope. The bracketed fields are deliberate evidence slots. They do not describe a chosen domain or working product. Once identity is resolved, replace the sentence with concrete nouns and verbs, then remove any feature that does not contribute to that primary outcome.')

s=base(F[2],'Show the path from input to action.','Illustrative architecture only. Every component below is proposed, not implemented.',6)
for x,label,body in [(0.58,'Interface','Accessible task input'),(3.72,'API / rules','Validate and process'),(6.86,'Persistence','Store required state'),(10.0,'Result','Explain next action')]:
    panel(s,x,3.2,2.72,2.55,label,body,'PROPOSED')
for x in [3.35,6.49,9.63]: arrow(s,x,4.43,.28)
txt(s,'Choose stack after domain review. Add trust boundaries, data rights and failure paths.',.6,6.24,12,.42,18,MUTED)
notes(s,55,'This diagram is a generic planning model, not the architecture of a completed SIH2171 prototype. When the domain is verified, choose the simplest stack that meets the real data, security and deployment constraints. Trace one input through validation, core processing, required persistence and an explainable result. Identify trust boundaries, external services and any simulated component. Explain the hardest implemented mechanism in plain language and make its correctness test visible. If persistence is unnecessary, remove that box. If the problem requires hardware, regulated data or offline operation, redesign the path rather than forcing it into this diagram.')

s=base(F[2],'Demonstrate a complete task.','A screen is not evidence of a workflow. Capture the actual delivered build.',7)
box(s,.58,3.02,7.75,3.4,WHITE,FOREST)
txt(s,'DEMO CAPTURE SLOT',.85,3.35,7.2,.4,16,CAUTION,True)
txt(s,'No prototype screenshot\nor video available.',.85,4.13,7.0,1.14,31,INK,True)
txt(s,'Replace with a real workflow capture after functional validation.',.85,5.62,7.0,.55,18,MUTED)
txt(s,'01  Start from known data\n02  Complete the core action\n03  Verify stored result\n04  Explain the outcome',8.89,3.35,3.85,2.78,22,INK,True)
notes(s,90,'Reserve this slot for a working demonstration, not a description of future screens. Before rehearsal, prepare deterministic seed data and record the build version. Begin at the known start state, perform the core action and verify the result independently where possible. Explain what the user sees and why the output matters. Keep the narration synchronized with the visible application. This blueprint currently contains no prototype capture and must not be presented as if it does. After implementation, replace the slot with the live flow or an explicitly labeled recording of the same tested build. Allow roughly forty seconds of this slot for the operator to perform the task.')

s=base(F[2],'Recovery is part of the product.','Plan one safe failure and prove the user can finish afterward.',8)
for x,title,body in [(.58,'Invalid input','Specific error\nValues preserved'),(4.72,'Service unavailable','Clear status\nRetry or fallback'),(8.86,'Return to task','No duplicate action\nConsistent final state')]:
    panel(s,x,3.12,3.85,2.94,title,body,'TEST PLAN')
txt(s,'Fallback ladder: live demo → recorded build → labeled screenshot storyboard',.6,6.34,12,.35,18,FOREST,True)
notes(s,45,'A happy path alone cannot demonstrate reliability. Choose one invalid input and one unavailable dependency that can be exercised safely. Check that errors identify the problem, preserve useful input and offer a meaningful next step. On retry, verify that the system does not create duplicate state. All of these are planned checks, not passing results. For the presentation itself, keep a local recording of the tested build and a static storyboard, then test playback without internet. State which fallback is being used. A recording proves a prior run of that build; it does not prove live uptime or a real external integration.')

s=base(F[3],'Publish the result, not the target.','Domain-product acceptance budgets. Preparation-site measurements are recorded separately.',9)
for i,(label,target,status) in enumerate([('Core flow response','< 200 ms','NOT MEASURED'),('Domain product website','> 90 each category','NOT BUILT'),('Accessibility','WCAG 2.1 AA review','NOT COMPLETE'),('Judge persuasion','No calibrated model','NOT ESTIMABLE')]):
    y=3.05+i*.74; rule(s,.6,y,12.1)
    txt(s,label,.6,y+.18,4.6,.39,21,INK,True)
    txt(s,target,5.3,y+.18,3.8,.39,19,MUTED)
    txt(s,status,9.45,y+.19,3.1,.4,15,CAUTION,True)
notes(s,50,'Separate acceptance thresholds from observations. The brief requests sub-two-hundred-millisecond core flows, Lighthouse scores above ninety in every category and WCAG 2.1 AA accessibility. These are targets. This slide does not claim that a domain product has passed them. The separate preparation site has local Lighthouse results in Docs/verification.md; those do not establish solution performance. Replace each open status with the actual report, environment, sample size and test scope after validation. Lighthouse accessibility is not a complete conformance assessment. Include manual keyboard and screen-reader review as appropriate. Judge persuasion probability cannot be estimated from the current research; a checklist score must not be relabeled as a chance of winning. Report failures and fixes as part of the engineering evidence.','User-requested budgets; Guardrails/; Benchmarks/; Wiki/pitch-strategy.md')

s=base(F[3],'Make operation credible.','Prototype simplicity first. Scale and security claims need demonstrated limits.',10)
panel(s,.58,3.06,3.85,2.99,'Run it','Start / reset / restore\nKnown dataset\nDeployment recipe','REQUIRED EVIDENCE')
panel(s,4.72,3.06,3.85,2.99,'Protect it','Validate inputs\nLeast-privilege access\nNo exposed secrets','REQUIRED EVIDENCE')
panel(s,8.86,3.06,3.85,2.99,'Bound it','Measured capacity\nDependency costs\nKnown failure limits','REQUIRED EVIDENCE')
txt(s,'No production-readiness, capacity or cost-saving claim is established.',.6,6.34,12,.34,18,CAUTION)
notes(s,45,'Explain how another team member can start, reset and recover the application from the README. Keep credentials outside version control and restrict access to the minimum required privileges. The final security scope depends on the actual domain and data classification, which are not yet known. A working localhost demonstration is not proof of production readiness. Capacity claims need load conditions, observed behavior and a stated limit. Cost estimates need a dated tariff source, a workload assumption and a support owner. Favor reversible choices and avoid services that add operational work without improving the primary task.')

s=base(F[3],'Six owners. Inspectable contributions.','Proposed work allocation only; experience and accepted commitments are not asserted.',11)
team=[('Francis','Integration + pitch'),('Gopreet','API + persistence'),('Hiranmayi','UX + accessibility'),('Varun','Domain implementation'),('Koushaik','Testing + reliability'),('Niharika','Research + narrative')]
for i,(name,role) in enumerate(team):
    x=.6+(i%3)*4.16; y=3.08+(i//3)*1.27
    rule(s,x,y,3.85); txt(s,name,x,y+.16,3.8,.42,25,INK,True); txt(s,role,x,y+.7,3.8,.37,18,MUTED)
box(s,.58,5.95,12.13,.59,LIME)
txt(s,'VERIFY → BUILD → TEST → REHEARSE  /  Internal target: 11 September 2026',.79,6.1,11.7,.32,17,INK,True)
notes(s,45,'These names and the university affiliation were supplied in the brief. The responsibilities are proposed ownership, not claims about existing skills or consented commitments. Francis integrates the work and technical explanation; Gopreet owns the API and persistence; Hiranmayi owns UX and accessibility; Varun owns the domain mechanism; Koushaik owns reliability evidence; and Niharika owns research and narrative. Each role should be backed by a real artifact before rehearsal. The internal target is eleven September, with the actual institutional cutoff time still unknown. Verification remains the first dependency. If it arrives late, remove optional features before cutting correctness or evidence.','Docs/team-plan.md; user brief')

s=base(F[4],'Impact is an outcome, not a user count.','A pilot must connect adoption to a meaningful change in the verified domain.',12,True)
txt(s,'Outcome change =',.65,3.3,11.8,.6,30,LIME,True)
txt(s,'[baseline result] − [pilot result]',.65,4.2,11.8,.92,42,PAPER,True)
txt(s,'Same task. Comparable population. Defined time window.\nReport uncertainty and unintended effects.',.65,5.65,11.7,.8,22,PAPER)
notes(s,45,'Once the domain is verified, choose an outcome that corresponds to the stated pain: time spent, task success, cost or another defensible measure. Define the baseline before the intervention, keep the task and population comparable and report the observation period. The subtraction shown here is a template for an outcome where lower is better; reverse or replace it when the domain requires a different metric. Do not treat account registrations as demonstrated benefit. Report uncertainty, accessibility differences and unintended consequences. No pilot population, observed improvement or social-impact percentage is available at this stage.')

s=base(F[4],'Design a pilot that can be sustained.','A proposed model needs a real adopter, realistic costs and a decision threshold.',13)
panel(s,.58,3.08,3.85,2.95,'Adoption','Named user organization\nWilling operator\nAccess and data rights','UNVERIFIED')
panel(s,4.72,3.08,3.85,2.95,'Economics','Unit operating cost\nSupport responsibility\nFunding assumption','NOT CALCULATED')
panel(s,8.86,3.08,3.85,2.95,'Pilot decision','Outcome threshold\nRisk review\nContinue / revise / stop','NOT DEFINED')
notes(s,40,'A sustainable proposal describes who will operate the solution after the hackathon and why they would adopt it. Identify a potential organization through evidence; do not invent a partnership or letter of support. Estimate operating cost per relevant task, including support and necessary external services. State whether funding would come from procurement, a subscription, a grant or another route only when the domain supports that model. Before a pilot begins, define a threshold for success and a rule for revising or stopping it. These are unfilled design tasks today, not a business case that has been validated.')

s=base(F[5],'Ground the work. Label the inference.','Historical competition evidence is useful preparation; domain validation is still missing.',14)
refs=[('01','Historical SIH criteria','Official 2024 indexed guideline; current applicability unverified.'),('02','Documented winner practice','Sponsor account: Solar Masters, 2024. Awards do not prove causality.'),('03','Domain + competitor sources','Required after SIH2171 identity is verified; no substitutes inserted.')]
for i,(n,title,body) in enumerate(refs):
    y=3.04+i*.98; txt(s,n,.6,y,.7,.45,22,FOREST,True); txt(s,title,1.42,y,10.9,.42,25,INK,True); txt(s,body,1.42,y+.49,10.8,.36,18,MUTED)
txt(s,'Source URLs and retrieval limitations are in speaker notes and Raw/competition-sources.json.',.6,6.4,12,.26,14,MUTED)
notes(s,50,'The competition research uses an official historical guideline, institutional award reports, sponsor records and labeled participant anecdotes. It is a purposive sample, not a census of winners. The official SIH PDF was indexed but direct retrieval returned a forbidden response, so current-year applicability remains unverified. The MathWorks winner interview provides an attributed account of prototype and presentation work; it does not establish that those practices caused winning. Technical literature, comparator documentation and problem-domain statistics still depend on the missing statement identity. Keep the Raw register with attribution and access limitations. Update this slide with the actual domain sources before any solution pitch.','https://sih.gov.in/letters/Guidelines-College-SPOC.pdf\nhttps://blogs.mathworks.com/student-lounge/2025/06/13/innovation-meets-excellence-solar-masters-winning-journey-at-smart-india-hackathon-2024/\nhttps://www.mathworks.com/academia/students/competitions/hackathons/winners.html\nRaw/competition-sources.json; Wiki/competition-intelligence.md')

s=base(F[5],'Earn the next decision\nwith evidence.','The completed entry should close on a bounded pilot ask. This blueprint cannot yet do that.',15,True)
box(s,.6,3.54,12.1,1.48,LIME)
txt(s,'Verify SIH2171. Build the core task. Prove the result.',.86,3.88,11.58,.8,29,INK,True)
txt(s,'OPEN DEPENDENCIES',.62,5.54,3.7,.4,14,LIME,True)
txt(s,'Official problem  ·  Domain prototype  ·  Measured validation  ·  Team rehearsal',.62,6.04,12,.45,20,PAPER)
notes(s,45,'The eventual closing should ask for one concrete next decision, such as a bounded pilot evaluated against transparent measures. That ask must follow demonstrated functionality and a verified problem. Today the actionable next step is to resolve SIH2171’s identity, implement the core task and collect real validation. Do not remove the not-submission-ready label until the unresolved fields are replaced and the artifacts are checked. Keep the expanded presentation separate from the supplied six-slide submission format. The planned slots total eleven minutes and forty seconds, with twenty seconds allowed for transitions; this has not been rehearsed.')

assert len(prs.slides)==15
for slide in prs.slides:
    assert any('NOT SUBMISSION READY' in sh.text for sh in slide.shapes if sh.has_text_frame)
    for sh in slide.shapes:
        assert sh.left >= 0 and sh.top >= 0
        assert sh.left+sh.width <= prs.slide_width+2
        assert sh.top+sh.height <= prs.slide_height+2
OUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUT)
print(f'Created {OUT}: 15 slides, editable shapes, speaker notes, 16:9.')
