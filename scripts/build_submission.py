#!/usr/bin/env python3
"""Preserve the supplied six-slide design and pointers; add honest preparation fields."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[1]
p=Presentation(R/'Raw/provided-template.pptx')
# The seventh slide is instructions, not a submitted content slide.
r=p.slides._sldIdLst[-1];p.part.drop_rel(r.rId);p.slides._sldIdLst.remove(r)
def text(shape,value,size):
 shape.text=value
 for para in shape.text_frame.paragraphs:
  para.font.name='Arial';para.font.size=Pt(size);para.font.color.rgb=RGBColor.from_string('15211F')
 shape.text_frame.word_wrap=True
p.slides[0].shapes[4].text='SMART INDIA HACKATHON 2026'
text(p.slides[0].shapes[5], '\nProblem Statement ID – SIH2171 (unverified)\nProblem Statement Title – pending official source\nTheme – unverified\nPS Category – unverified\nTeam ID – not supplied\nTeam Name (Registered on portal) – not supplied',20)
# Keep every provided idea-details pointer; reserve the remaining space for explicit status.
additions=[None,
 'Proposed Solution (Describe your Idea/Solution/Prototype)\nPENDING: verify the official problem before defining a solution.\n\nDetailed explanation of the proposed solution\nHow it addresses the problem\nInnovation and uniqueness of the solution\nPENDING: requirements, mechanism and baseline comparison.',
 'Technologies to be used (e.g. programming languages, frameworks, hardware)\nPROVISIONAL: React / TypeScript + Node; database only if required.\n\nMethodology and process for implementation (Flow Charts/Images/working prototype)\nVerify statement → freeze flows → build → test → demonstrate.\nNo domain prototype has been implemented.',
 'Analysis of the feasibility of the idea\nPENDING: official constraints and data availability.\nPotential challenges and risks\nUnverified ID; unknown domain scope; three-day delivery window.\nStrategies for overcoming these challenges\nResolve source identity; one vertical slice; deterministic demo fallback.',
 'Potential impact on the target audience\nPENDING: verified beneficiary, baseline and measurable outcome.\nBenefits of the solution (social, economic, environmental, etc.)\nPENDING: domain-specific comparison and pilot evidence.\nNo impact percentages or achieved benefits are claimed.',
 'Details / Links of the reference and research work\nOfficial catalogue: https://sih.gov.in/sih2026PS (identity unresolved)\nHistorical criteria: https://sih.gov.in/letters/Guidelines-College-SPOC.pdf\nWinner evidence: https://www.mathworks.com/academia/students/competitions/hackathons/winners.html\nDomain literature: pending verified problem.\nSource register and access limits: Raw/competition-sources.json'
]
for i,s in enumerate(p.slides):
 if i:
  body=s.shapes[2];body.top=Inches(1.85);body.left=Inches(.7);body.width=Inches(11.9);body.height=Inches(4.55);text(body,additions[i],22 if i<5 else 19)
  text(s.shapes[5],'TEAM NAME\nPENDING',10)
 notice=s.shapes.add_textbox(Inches(.7),Inches(6.35),Inches(11.9),Inches(.42))
 text(notice,'PREPARATION BLUEPRINT — NOT SUBMISSION READY',15)
 for para in notice.text_frame.paragraphs:para.font.bold=True;para.font.color.rgb=RGBColor.from_string('9F271D')
 s.notes_slide.notes_text_frame.text='Derived from the user-supplied 2025 template. Six content slides and original heading families/pointers preserved. 2026 applicability, SIH2171 identity and registered details require verification. This is an editable blueprint, not a completed idea submission. See Docs/template-audit.md.'
p.core_properties.title='SIH2171 | Six-slide submission blueprint | NOT SUBMISSION READY'
p.save(R/'Docs/submission-blueprint.pptx')
print('Saved six-slide supplied-template blueprint')
